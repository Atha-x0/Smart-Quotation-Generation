from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook


PARSER_NAME = "machine_details_connected_load_parser"


def clean_text(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def normalize_header(value: Any) -> str:
    text = clean_text(value).lower()
    text = text.replace(".", "")
    text = text.replace("_", " ")
    text = text.replace("-", " ")
    text = " ".join(text.split())
    return text


def to_number(value: Any) -> float:
    if value is None:
        return 0.0

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()

    if not text:
        return 0.0

    text = text.replace(",", "")
    text = text.replace("kw", "")
    text = text.replace("KW", "")
    text = text.replace("hp", "")
    text = text.replace("HP", "")
    text = text.replace("rpm", "")
    text = text.replace("RPM", "")
    text = text.strip()

    try:
        return float(text)
    except Exception:
        return 0.0


def is_section_row(row_values: List[Any]) -> bool:
    non_empty_values = [clean_text(value) for value in row_values if clean_text(value)]

    if len(non_empty_values) == 0:
        return True

    if len(non_empty_values) == 1:
        return True

    first_value = to_number(row_values[0]) if len(row_values) > 0 else 0
    kw_value = to_number(row_values[2]) if len(row_values) > 2 else 0

    if first_value <= 0 and kw_value <= 0:
        return True

    return False


def find_header_row(worksheet) -> Tuple[Optional[int], Dict[str, int]]:
    header_keywords = {
        "serial": ["sr no", "sr", "s no", "sno", "serial", "serial no"],
        "equipment": ["name", "equipment", "machine", "description", "load"],
        "kw": ["kw", "rating kw", "motor kw", "load kw", "connected kw"],
        "hp": ["hp", "rating hp", "motor hp"],
        "rpm": ["rpm", "speed"],
        "quantity": ["qty", "quantity", "nos", "no of"],
        "location": ["location", "area", "section", "department"],
        "hours": ["hours", "running hours", "operating hours"],
        "diversity": ["diversity", "diversity factor", "df"],
    }

    max_rows_to_scan = min(40, worksheet.max_row)

    for row_number in range(1, max_rows_to_scan + 1):
        row = list(
            worksheet.iter_rows(
                min_row=row_number,
                max_row=row_number,
                values_only=True,
            )
        )[0]

        normalized_cells = [normalize_header(cell) for cell in row]
        column_map: Dict[str, int] = {}

        for column_index, header_text in enumerate(normalized_cells):
            if not header_text:
                continue

            for field_name, keywords in header_keywords.items():
                if field_name in column_map:
                    continue

                if header_text in keywords:
                    column_map[field_name] = column_index

        has_equipment = "equipment" in column_map
        has_kw = "kw" in column_map

        if has_equipment and has_kw:
            return row_number, column_map

    return None, {}


def get_cell(row_values: List[Any], column_map: Dict[str, int], field_name: str) -> Any:
    column_index = column_map.get(field_name)

    if column_index is None:
        return None

    if column_index >= len(row_values):
        return None

    return row_values[column_index]


def build_row_warnings(
    equipment_name: str,
    rating_kw: float,
    rating_hp: float,
    quantity: float,
    diversity_factor: float,
) -> List[str]:
    warnings: List[str] = []

    if not equipment_name:
        warnings.append("Equipment name is missing.")

    if rating_kw <= 0 and rating_hp <= 0:
        warnings.append("Both kW and HP rating are missing or zero.")

    if quantity <= 0:
        warnings.append("Quantity was missing or zero. Default quantity 1 was used.")

    if diversity_factor <= 0:
        warnings.append("Diversity factor was missing or zero. Default diversity factor 1 was used.")

    return warnings


def calculate_row_confidence(
    equipment_name: str,
    rating_kw: float,
    rating_hp: float,
    quantity: float,
    diversity_factor: float,
) -> int:
    score = 100

    if not equipment_name:
        score -= 30

    if rating_kw <= 0 and rating_hp <= 0:
        score -= 40

    if quantity <= 0:
        score -= 10

    if diversity_factor <= 0:
        score -= 10

    return max(0, min(score, 100))


def build_connected_load_row(
    source_file: str,
    sheet_name: str,
    row_number: int,
    header_row_number: int,
    row_values: List[Any],
    column_map: Dict[str, int],
    current_section: str,
) -> Optional[Dict[str, Any]]:
    equipment_name = clean_text(get_cell(row_values, column_map, "equipment"))

    if not equipment_name:
        return None

    sr_no = to_number(get_cell(row_values, column_map, "serial"))
    rating_kw = to_number(get_cell(row_values, column_map, "kw"))
    rating_hp = to_number(get_cell(row_values, column_map, "hp"))
    rpm = to_number(get_cell(row_values, column_map, "rpm"))

    raw_quantity = to_number(get_cell(row_values, column_map, "quantity"))
    raw_diversity_factor = to_number(get_cell(row_values, column_map, "diversity"))
    raw_running_hours = to_number(get_cell(row_values, column_map, "hours"))

    location = clean_text(get_cell(row_values, column_map, "location"))

    if not location:
        location = current_section

    row_warnings = build_row_warnings(
        equipment_name=equipment_name,
        rating_kw=rating_kw,
        rating_hp=rating_hp,
        quantity=raw_quantity,
        diversity_factor=raw_diversity_factor,
    )

    row_confidence = calculate_row_confidence(
        equipment_name=equipment_name,
        rating_kw=rating_kw,
        rating_hp=rating_hp,
        quantity=raw_quantity,
        diversity_factor=raw_diversity_factor,
    )

    quantity = raw_quantity if raw_quantity > 0 else 1
    diversity_factor = raw_diversity_factor if raw_diversity_factor > 0 else 1
    running_hours = raw_running_hours if raw_running_hours > 0 else 0

    if rating_kw <= 0 and rating_hp > 0:
        rating_kw = rating_hp * 0.746

    connected_kw = rating_kw * quantity * diversity_factor

    if rating_kw <= 0 and rating_hp <= 0:
        return None

    return {
        "sourceFile": source_file,
        "sheetName": sheet_name,
        "rowNumber": row_number,
        "headerRowNumber": header_row_number,
        "equipmentName": equipment_name,
        "location": location,
        "quantity": round(quantity, 2),
        "ratingKW": round(rating_kw, 3),
        "ratingHP": round(rating_hp, 3),
        "runningHours": round(running_hours, 2),
        "diversityFactor": round(diversity_factor, 3),
        "connectedKW": round(connected_kw, 3),
        "rpm": round(rpm, 2),
        "serialNumber": int(sr_no) if sr_no > 0 else "",
        "parser": PARSER_NAME,
        "extractionConfidence": row_confidence,
        "extractionWarnings": row_warnings,
    }


def parse_sheet(
    worksheet,
    source_file: str,
    sheet_name: str,
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    header_row_number, column_map = find_header_row(worksheet)

    if header_row_number is None:
        raise ValueError("No header row found in Excel file.")

    connected_loads: List[Dict[str, Any]] = []
    current_section = ""
    skipped_section_rows = 0
    skipped_blank_rows = 0
    skipped_invalid_rows = 0

    for row_number in range(header_row_number + 1, worksheet.max_row + 1):
        row_values = list(
            worksheet.iter_rows(
                min_row=row_number,
                max_row=row_number,
                values_only=True,
            )
        )[0]

        row_values_list = list(row_values)
        non_empty_values = [clean_text(value) for value in row_values_list if clean_text(value)]

        if len(non_empty_values) == 0:
            skipped_blank_rows += 1
            continue

        if is_section_row(row_values_list):
            section_name = ""

            for value in row_values_list:
                text = clean_text(value)
                if text and to_number(text) <= 0:
                    section_name = text
                    break

            if section_name:
                current_section = section_name
                skipped_section_rows += 1
            else:
                skipped_blank_rows += 1

            continue

        connected_load_row = build_connected_load_row(
            source_file=source_file,
            sheet_name=sheet_name,
            row_number=row_number,
            header_row_number=header_row_number,
            row_values=row_values_list,
            column_map=column_map,
            current_section=current_section,
        )

        if connected_load_row is not None:
            connected_loads.append(connected_load_row)
        else:
            skipped_invalid_rows += 1

    average_confidence = 0

    if connected_loads:
        average_confidence = round(
            sum(row.get("extractionConfidence", 0) for row in connected_loads)
            / len(connected_loads)
        )

    sheet_summary = {
        "sheetName": sheet_name,
        "headerRowNumber": header_row_number,
        "columnMap": column_map,
        "rowsExtracted": len(connected_loads),
        "skippedSectionRows": skipped_section_rows,
        "skippedBlankRows": skipped_blank_rows,
        "skippedInvalidRows": skipped_invalid_rows,
        "averageExtractionConfidence": average_confidence,
    }

    return connected_loads, sheet_summary


def parse_connected_load_file_with_summary(
    file_path: Path,
    source_file: str = "",
) -> Dict[str, Any]:
    workbook = load_workbook(file_path, data_only=True, read_only=True)
    all_connected_loads: List[Dict[str, Any]] = []
    sheet_summaries: List[Dict[str, Any]] = []
    parser_warnings: List[str] = []

    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]

            try:
                sheet_rows, sheet_summary = parse_sheet(
                    worksheet=worksheet,
                    source_file=source_file or Path(file_path).name,
                    sheet_name=sheet_name,
                )

                all_connected_loads.extend(sheet_rows)
                sheet_summaries.append(sheet_summary)

            except ValueError as error:
                parser_warnings.append(f"{sheet_name}: {error}")

        if not all_connected_loads:
            raise ValueError("No connected load rows could be extracted from Excel file.")

        total_connected_kw = round(
            sum(row.get("connectedKW", 0) for row in all_connected_loads),
            3,
        )

        average_confidence = round(
            sum(row.get("extractionConfidence", 0) for row in all_connected_loads)
            / len(all_connected_loads)
        )

        return {
            "connectedLoads": all_connected_loads,
            "summary": {
                "parser": PARSER_NAME,
                "sourceFile": source_file or Path(file_path).name,
                "sheetCount": len(workbook.sheetnames),
                "rowsExtracted": len(all_connected_loads),
                "totalConnectedKW": total_connected_kw,
                "averageExtractionConfidence": average_confidence,
                "sheetSummaries": sheet_summaries,
                "parserWarnings": parser_warnings,
            },
        }

    finally:
        workbook.close()


def parse_connected_load_file(file_path: Path, source_file: str = "") -> List[Dict[str, Any]]:
    result = parse_connected_load_file_with_summary(file_path, source_file)
    return result["connectedLoads"]


def parse_load_file(file_path: Path, source_file: str = "") -> List[Dict[str, Any]]:
    return parse_connected_load_file(file_path, source_file)


def parse_connected_loads(file_path: Path, source_file: str = "") -> List[Dict[str, Any]]:
    return parse_connected_load_file(file_path, source_file)