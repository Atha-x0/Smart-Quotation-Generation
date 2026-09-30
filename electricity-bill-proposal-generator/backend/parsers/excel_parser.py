from typing import Any, Dict, List
from openpyxl import load_workbook


def to_number(value: Any) -> float:
    if value is None or value == "":
        return 0

    if isinstance(value, (int, float)):
        return float(value)

    cleaned = str(value).replace(",", "").replace("₹", "").strip()

    try:
        return float(cleaned)
    except ValueError:
        return 0


def to_text(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def normalize_key(key: str) -> str:
    return (
        key.lower()
        .replace(" ", "")
        .replace("_", "")
        .replace(".", "")
        .replace("(", "")
        .replace(")", "")
        .replace("/", "")
    )


def get_value(row: Dict[str, Any], possible_keys: List[str]) -> Any:
    normalized_row = {}

    for key, value in row.items():
        normalized_row[normalize_key(key)] = value

    for key in possible_keys:
        normalized_key = normalize_key(key)

        if normalized_key in normalized_row:
            return normalized_row[normalized_key]

    return None


def parse_excel_file(file_path: str) -> List[Dict[str, Any]]:
    workbook = load_workbook(filename=file_path, data_only=True)
    worksheet = workbook.active

    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        raise ValueError("Excel file is empty.")

    headers = [to_text(cell) for cell in rows[0]]

    if not any(headers):
        raise ValueError("No header row found in Excel file.")

    data_rows: List[Dict[str, Any]] = []

    for row in rows[1:]:
        row_data: Dict[str, Any] = {}

        for index, header in enumerate(headers):
            if not header:
                continue

            row_data[header] = row[index] if index < len(row) else None

        if any(value not in [None, ""] for value in row_data.values()):
            data_rows.append(row_data)

    bills = []

    for index, row in enumerate(data_rows):
        consumer_name = to_text(
            get_value(
                row,
                [
                    "Consumer Name",
                    "Company Name",
                    "Client Name",
                    "Name",
                    "Customer Name",
                ],
            )
        )

        consumer_number = to_text(
            get_value(
                row,
                [
                    "Consumer No",
                    "Consumer Number",
                    "Account No",
                    "Account Number",
                    "Customer No",
                ],
            )
        )

        tariff_category = to_text(
            get_value(
                row,
                [
                    "Tariff",
                    "Tariff Category",
                    "Category",
                ],
            )
        )

        month = (
            to_text(
                get_value(
                    row,
                    [
                        "Month",
                        "Bill Month",
                        "Billing Month",
                        "Period",
                        "Date",
                    ],
                )
            )
            or f"Month {index + 1}"
        )

        contract_demand_kva = to_number(
            get_value(
                row,
                [
                    "Contract Demand",
                    "Contract Demand KVA",
                    "CD",
                    "CD KVA",
                    "Sanctioned Demand",
                    "Sanctioned Load",
                ],
            )
        )

        actual_demand_kva = to_number(
            get_value(
                row,
                [
                    "Actual Demand",
                    "Maximum Demand",
                    "MD",
                    "Recorded Demand",
                    "Actual Demand KVA",
                    "Max Demand",
                ],
            )
        )

        billing_demand_kva = to_number(
            get_value(
                row,
                [
                    "Billing Demand",
                    "Billing Demand KVA",
                    "Bill Demand",
                    "Billed Demand",
                ],
            )
        )

        kwh = to_number(
            get_value(
                row,
                [
                    "kWh",
                    "KWH",
                    "Units",
                    "Consumption",
                    "Energy Consumption",
                    "kWh Consumption",
                ],
            )
        )

        kvah = to_number(
            get_value(
                row,
                [
                    "kVAh",
                    "KVAH",
                    "kVAh Consumption",
                    "Apparent Energy",
                ],
            )
        )

        direct_power_factor = to_number(
            get_value(
                row,
                [
                    "PF",
                    "Power Factor",
                    "Avg PF",
                    "Average PF",
                ],
            )
        )

        if direct_power_factor > 0:
            power_factor = direct_power_factor
        elif kwh > 0 and kvah > 0:
            power_factor = kwh / kvah
        else:
            power_factor = 0

        demand_charge = to_number(
            get_value(
                row,
                [
                    "Demand Charge",
                    "Demand Charges",
                    "MD Charges",
                    "Fixed Charge",
                    "Fixed Charges",
                ],
            )
        )

        energy_charge = to_number(
            get_value(
                row,
                [
                    "Energy Charge",
                    "Energy Charges",
                    "Consumption Charge",
                    "Consumption Charges",
                    "Variable Charge",
                ],
            )
        )

        pf_penalty = to_number(
            get_value(
                row,
                [
                    "PF Penalty",
                    "Power Factor Penalty",
                    "Reactive Penalty",
                ],
            )
        )

        pf_incentive = to_number(
            get_value(
                row,
                [
                    "PF Incentive",
                    "Power Factor Incentive",
                ],
            )
        )

        tod_charges = to_number(
            get_value(
                row,
                [
                    "TOD",
                    "TOD Charges",
                    "Time of Day Charges",
                ],
            )
        )

        other_charges = to_number(
            get_value(
                row,
                [
                    "Other Charges",
                    "Other Charge",
                    "Taxes",
                    "Duty",
                    "Electricity Duty",
                    "FAC",
                ],
            )
        )

        total_bill_amount = to_number(
            get_value(
                row,
                [
                    "Total Bill",
                    "Total Bill Amount",
                    "Bill Amount",
                    "Net Amount",
                    "Total Amount",
                    "Amount Payable",
                ],
            )
        )

        bill = {
            "month": month,
            "consumerName": consumer_name,
            "consumerNumber": consumer_number,
            "tariffCategory": tariff_category,
            "contractDemandKVA": contract_demand_kva,
            "actualDemandKVA": actual_demand_kva,
            "billingDemandKVA": billing_demand_kva,
            "kWh": kwh,
            "kVAh": kvah,
            "powerFactor": power_factor,
            "demandCharge": demand_charge,
            "energyCharge": energy_charge,
            "pfPenalty": pf_penalty,
            "pfIncentive": pf_incentive,
            "todCharges": tod_charges,
            "otherCharges": other_charges,
            "totalBillAmount": total_bill_amount,
        }

        if (
            bill["month"]
            or bill["kWh"] > 0
            or bill["kVAh"] > 0
            or bill["totalBillAmount"] > 0
        ):
            bills.append(bill)

    return bills