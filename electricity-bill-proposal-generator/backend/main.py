from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Dict, List, Tuple

from fastapi import FastAPI, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware

from parsers.pdf_parser import build_pdf_debug, parse_pdf_file

try:
    from parsers.pdf_parser import parse_bill_files
except Exception:
    parse_bill_files = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from parsers.load_parser import parse_connected_load_file
except Exception:
    try:
        from parsers.load_parser import parse_load_file as parse_connected_load_file
    except Exception:
        try:
            from parsers.load_parser import parse_connected_loads as parse_connected_load_file
        except Exception:
            parse_connected_load_file = None

try:
    from parsers.excel_parser import parse_excel_file
except Exception:
    parse_excel_file = None

try:
    from parsers.csv_parser import parse_csv_file
except Exception:
    parse_csv_file = None


app = FastAPI(title="Electricity Bill Proposal Generator API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def save_upload_to_temp(upload_file: UploadFile) -> Path:
    suffix = Path(upload_file.filename or "uploaded_file").suffix

    with NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
        temp_file.write(upload_file.file.read())
        return Path(temp_file.name)


def safe_remove(path: Path) -> None:
    try:
        if path.exists():
            path.unlink()
    except Exception:
        pass


def get_file_extension(filename: str) -> str:
    return Path(filename or "").suffix.lower()


# =========================================================
# BILL INPUT SUPPORT
# PDF bills are already supported.
# Image bills will work after OCR support is enabled in parsers/pdf_parser.py
# and OCR dependencies are installed.
# =========================================================
PDF_BILL_EXTENSIONS = [".pdf"]
IMAGE_BILL_EXTENSIONS = [".jpg", ".jpeg", ".png", ".webp"]
BILL_FILE_EXTENSIONS = PDF_BILL_EXTENSIONS + IMAGE_BILL_EXTENSIONS


def is_bill_file_extension(extension: str) -> bool:
    return extension in BILL_FILE_EXTENSIONS


def classify_file(filename: str) -> str:
    extension = get_file_extension(filename)

    if extension in BILL_FILE_EXTENSIONS:
        return "BILL_FILE"

    if extension in [".xlsx", ".xls"]:
        name = filename.lower()

        if "connected" in name or "load" in name:
            return "CONNECTED_LOAD_FILE"

        if "machine" in name or "equipment" in name:
            return "MACHINE_DETAILS_FILE"

        return "EXCEL_FILE"

    if extension == ".csv":
        return "CSV_FILE"

    return "UNKNOWN_FILE"


def normalize_bill_result(result: Any) -> List[Dict[str, Any]]:
    if isinstance(result, list):
        return [item for item in result if isinstance(item, dict)]

    if isinstance(result, dict):
        return [result]

    return []


def run_connected_load_parser(
    temp_path: Path,
    filename: str,
) -> Tuple[List[Dict[str, Any]], List[str]]:
    warnings: List[str] = []

    if parse_connected_load_file is None:
        warnings.append("Connected load parser is not available.")
        return [], warnings

    try:
        try:
            result = parse_connected_load_file(temp_path, filename)
        except TypeError:
            result = parse_connected_load_file(temp_path)

        if isinstance(result, list):
            rows = [item for item in result if isinstance(item, dict)]
            return rows, warnings

        warnings.append("Connected load parser returned unsupported output format.")
        return [], warnings

    except Exception as error:
        warnings.append(str(error))
        return [], warnings


def run_excel_bill_parser(
    temp_path: Path,
    filename: str,
) -> Tuple[List[Dict[str, Any]], List[str]]:
    warnings: List[str] = []

    if parse_excel_file is None:
        warnings.append("Excel bill parser is not available.")
        return [], warnings

    try:
        try:
            result = parse_excel_file(temp_path, filename)
        except TypeError:
            result = parse_excel_file(temp_path)

        return normalize_bill_result(result), warnings

    except Exception as error:
        warnings.append(str(error))
        return [], warnings


def run_csv_parser(
    temp_path: Path,
    filename: str,
) -> Tuple[List[Dict[str, Any]], List[str]]:
    warnings: List[str] = []

    if parse_csv_file is None:
        warnings.append("CSV parser is not available.")
        return [], warnings

    try:
        try:
            result = parse_csv_file(temp_path, filename)
        except TypeError:
            result = parse_csv_file(temp_path)

        return normalize_bill_result(result), warnings

    except Exception as error:
        warnings.append(str(error))
        return [], warnings


def parse_bill_safely(temp_path: Path) -> List[Dict[str, Any]]:
    # parse_pdf_file now acts as the common bill parser:
    # - Digital PDF bills
    # - Scanned PDF bills, if OCR support exists in pdf_parser.py
    # - Image bill files, if OCR support exists in pdf_parser.py
    result = parse_pdf_file(temp_path)
    return normalize_bill_result(result)

def parse_multiple_bill_files_safely(temp_paths: List[Path]) -> List[Dict[str, Any]]:
    """
    Parse multiple bill files together.

    Important use case:
    - Front-side bill image + back-side bill image should become one bill row.
    - Digital PDF bills should continue to work normally.
    """
    if parse_bill_files is None:
        all_bills: List[Dict[str, Any]] = []

        for temp_path in temp_paths:
            all_bills.extend(parse_bill_safely(temp_path))

        return all_bills

    result = parse_bill_files([str(path) for path in temp_paths])
    return normalize_bill_result(result)

def parse_connected_load_safely(temp_path: Path, filename: str) -> List[Dict[str, Any]]:
    rows, _warnings = run_connected_load_parser(temp_path, filename)
    return rows


def parse_excel_safely(temp_path: Path, filename: str) -> List[Dict[str, Any]]:
    rows, _warnings = run_excel_bill_parser(temp_path, filename)
    return rows


def parse_csv_safely(temp_path: Path, filename: str) -> List[Dict[str, Any]]:
    rows, _warnings = run_csv_parser(temp_path, filename)
    return rows


def get_excel_workbook_debug(temp_path: Path) -> Dict[str, Any]:
    if load_workbook is None:
        return {
            "sheetNames": [],
            "sheetCount": 0,
            "firstSheetPreview": [],
            "error": "openpyxl is not available.",
        }

    try:
        workbook = load_workbook(temp_path, data_only=True, read_only=True)
        sheet_names = workbook.sheetnames

        first_sheet_preview: List[List[str]] = []

        if sheet_names:
            worksheet = workbook[sheet_names[0]]

            for row in worksheet.iter_rows(min_row=1, max_row=30, values_only=True):
                clean_row = []
                for cell in row:
                    if cell is None:
                        clean_row.append("")
                    else:
                        clean_row.append(str(cell))
                first_sheet_preview.append(clean_row)

        workbook.close()

        return {
            "sheetNames": sheet_names,
            "sheetCount": len(sheet_names),
            "firstSheetPreview": first_sheet_preview,
            "error": "",
        }

    except Exception as error:
        return {
            "sheetNames": [],
            "sheetCount": 0,
            "firstSheetPreview": [],
            "error": str(error),
        }


def calculate_total_connected_kw(connected_loads: List[Dict[str, Any]]) -> float:
    total = 0.0

    for row in connected_loads:
        value = row.get("connectedKW", 0)

        try:
            total += float(value or 0)
        except Exception:
            pass

    return total


def get_excel_detected_type(
    filename: str,
    connected_loads: List[Dict[str, Any]],
    bills: List[Dict[str, Any]],
) -> str:
    if connected_loads:
        return "CONNECTED_LOAD_FILE"

    if bills:
        return "EXCEL_BILL_FILE"

    return classify_file(filename)


@app.get("/")
def root() -> Dict[str, str]:
    return {
        "status": "running",
        "message": "Electricity Bill Proposal Generator backend is running.",
    }


@app.post("/api/debug-pdf")
async def debug_pdf(file: UploadFile = File(...)) -> Dict[str, Any]:
    temp_path = save_upload_to_temp(file)

    try:
        extension = get_file_extension(file.filename or "")

        if extension == ".pdf":
            debug_data = build_pdf_debug(temp_path)
        else:
            debug_data = {
                "detectedFormat": "IMAGE_BILL_INPUT",
                "lineCount": 0,
                "firstLines": [],
                "lines": [],
                "textPreview": "Image bill input. OCR text preview is available only if OCR is enabled in pdf_parser.py.",
            }

        parsed_bills = parse_bill_safely(temp_path)

        first_bill = parsed_bills[0] if parsed_bills else {}

        lines = debug_data.get("lines", [])
        text_preview = debug_data.get("textPreview", "")

        if not text_preview and lines:
            text_preview = "\n".join(lines[:150])

        return {
            "fileName": file.filename,
            "detectedFormat": debug_data.get(
                "detectedFormat",
                first_bill.get("detectedBillFormat", "UNKNOWN"),
            ),
            "lineCount": debug_data.get("lineCount", len(lines)),
            "firstLines": debug_data.get("firstLines", lines[:150]),
            "textPreview": text_preview,
            "parserResult": first_bill,
            "parsedBills": parsed_bills,
            "parserSummary": {
                "month": first_bill.get("month", ""),
                "contractDemandKVA": first_bill.get("contractDemandKVA", 0),
                "actualDemandKVA": first_bill.get("actualDemandKVA", 0),
                "billingDemandKVA": first_bill.get("billingDemandKVA", 0),
                "kWh": first_bill.get("kWh", 0),
                "kVAh": first_bill.get("kVAh", 0),
                "powerFactor": first_bill.get("powerFactor", 0),
                "demandCharge": first_bill.get("demandCharge", 0),
                "energyCharge": first_bill.get("energyCharge", 0),
                "totalBillAmount": first_bill.get("totalBillAmount", 0),
                "detectedBillFormat": first_bill.get("detectedBillFormat", ""),
                "parser": first_bill.get("parser", ""),
                "extractionConfidence": first_bill.get("extractionConfidence", 0),
                "extractionWarnings": first_bill.get("extractionWarnings", []),
            },
            "rawDebug": debug_data,
        }

    except Exception as error:
        return {
            "fileName": file.filename,
            "detectedFormat": "ERROR",
            "lineCount": 0,
            "firstLines": [],
            "textPreview": "",
            "parserResult": {},
            "parsedBills": [],
            "parserSummary": {
                "month": "",
                "contractDemandKVA": 0,
                "actualDemandKVA": 0,
                "billingDemandKVA": 0,
                "kWh": 0,
                "kVAh": 0,
                "powerFactor": 0,
                "demandCharge": 0,
                "energyCharge": 0,
                "totalBillAmount": 0,
                "detectedBillFormat": "ERROR",
                "parser": "",
                "extractionConfidence": 0,
                "extractionWarnings": [str(error)],
            },
            "rawDebug": {
                "error": str(error),
            },
        }

    finally:
        safe_remove(temp_path)


@app.post("/api/debug-excel")
async def debug_excel(file: UploadFile = File(...)) -> Dict[str, Any]:
    temp_path = save_upload_to_temp(file)

    try:
        extension = get_file_extension(file.filename or "")

        if extension not in [".xlsx", ".xls"]:
            return {
                "fileName": file.filename,
                "detectedExcelType": "UNSUPPORTED_FILE",
                "sheetNames": [],
                "sheetCount": 0,
                "connectedLoadRows": 0,
                "billRows": 0,
                "totalConnectedKW": 0,
                "firstConnectedLoadRows": [],
                "firstBillRows": [],
                "firstSheetPreview": [],
                "warnings": ["Please upload an Excel file with .xlsx or .xls extension."],
                "rawDebug": {},
            }

        workbook_debug = get_excel_workbook_debug(temp_path)

        connected_loads, connected_load_warnings = run_connected_load_parser(
            temp_path,
            file.filename or "",
        )

        bills, excel_bill_warnings = run_excel_bill_parser(
            temp_path,
            file.filename or "",
        )

        detected_excel_type = get_excel_detected_type(
            file.filename or "",
            connected_loads,
            bills,
        )

        warnings: List[str] = []

        if workbook_debug.get("error"):
            warnings.append(str(workbook_debug.get("error")))

        warnings.extend(connected_load_warnings)
        warnings.extend(excel_bill_warnings)

        if not connected_loads and not bills:
            warnings.append(
                "No connected load rows or bill rows could be extracted from this Excel file."
            )

        return {
            "fileName": file.filename,
            "detectedExcelType": detected_excel_type,
            "sheetNames": workbook_debug.get("sheetNames", []),
            "sheetCount": workbook_debug.get("sheetCount", 0),
            "connectedLoadRows": len(connected_loads),
            "billRows": len(bills),
            "totalConnectedKW": round(calculate_total_connected_kw(connected_loads), 2),
            "firstConnectedLoadRows": connected_loads[:50],
            "firstBillRows": bills[:50],
            "firstSheetPreview": workbook_debug.get("firstSheetPreview", []),
            "warnings": list(dict.fromkeys([warning for warning in warnings if warning])),
            "rawDebug": {
                "workbookDebug": workbook_debug,
                "connectedLoadParserWarnings": connected_load_warnings,
                "excelBillParserWarnings": excel_bill_warnings,
                "connectedLoads": connected_loads,
                "bills": bills,
            },
        }

    except Exception as error:
        workbook_debug = get_excel_workbook_debug(temp_path)

        return {
            "fileName": file.filename,
            "detectedExcelType": classify_file(file.filename or ""),
            "sheetNames": workbook_debug.get("sheetNames", []),
            "sheetCount": workbook_debug.get("sheetCount", 0),
            "connectedLoadRows": 0,
            "billRows": 0,
            "totalConnectedKW": 0,
            "firstConnectedLoadRows": [],
            "firstBillRows": [],
            "firstSheetPreview": workbook_debug.get("firstSheetPreview", []),
            "warnings": [str(error)],
            "rawDebug": {
                "error": str(error),
                "workbookDebug": workbook_debug,
            },
        }

    finally:
        safe_remove(temp_path)


@app.post("/api/extract")
async def extract_file(file: UploadFile = File(...)) -> Dict[str, Any]:
    temp_path = save_upload_to_temp(file)

    try:
        file_type = classify_file(file.filename or "")
        extension = get_file_extension(file.filename or "")

        if is_bill_file_extension(extension):
            bills = parse_bill_safely(temp_path)

            return {
                "status": "success",
                "fileName": file.filename,
                "fileType": "BILL_FILE",
                "bills": bills,
                "connectedLoads": [],
                "message": f"Extracted {len(bills)} bill row(s).",
            }

        if extension in [".xlsx", ".xls"]:
            connected_loads = parse_connected_load_safely(
                temp_path,
                file.filename or "",
            )

            if connected_loads:
                return {
                    "status": "success",
                    "fileName": file.filename,
                    "fileType": "CONNECTED_LOAD_FILE",
                    "bills": [],
                    "connectedLoads": connected_loads,
                    "message": f"Extracted {len(connected_loads)} connected load row(s).",
                }

            bills = parse_excel_safely(temp_path, file.filename or "")

            if bills:
                return {
                    "status": "success",
                    "fileName": file.filename,
                    "fileType": file_type,
                    "bills": bills,
                    "connectedLoads": [],
                    "message": f"Extracted {len(bills)} Excel bill row(s).",
                }

            return {
                "status": "error",
                "fileName": file.filename,
                "fileType": file_type,
                "bills": [],
                "connectedLoads": [],
                "message": "No data could be extracted from this Excel file.",
            }

        if extension == ".csv":
            bills = parse_csv_safely(temp_path, file.filename or "")

            if bills:
                return {
                    "status": "success",
                    "fileName": file.filename,
                    "fileType": "CSV_FILE",
                    "bills": bills,
                    "connectedLoads": [],
                    "message": f"Extracted {len(bills)} CSV bill row(s).",
                }

            return {
                "status": "error",
                "fileName": file.filename,
                "fileType": "CSV_FILE",
                "bills": [],
                "connectedLoads": [],
                "message": "No data could be extracted from this CSV file.",
            }

        return {
            "status": "error",
            "fileName": file.filename,
            "fileType": "UNKNOWN_FILE",
            "bills": [],
            "connectedLoads": [],
            "message": "Unsupported file type.",
        }

    except Exception as error:
        return {
            "status": "error",
            "fileName": file.filename,
            "fileType": classify_file(file.filename or ""),
            "bills": [],
            "connectedLoads": [],
            "message": str(error),
        }

    finally:
        safe_remove(temp_path)

def attach_source_file_to_bills(
    bills: List[Dict[str, Any]],
    filename: str,
) -> List[Dict[str, Any]]:
    """
    Attach the original uploaded filename to every parsed bill row.

    Reason:
    Parser receives a temporary file path like tmpabc123.pdf.
    Without this helper, dashboard shows the temp filename instead of
    the actual uploaded filename.
    """
    clean_bills: List[Dict[str, Any]] = []

    for bill in bills:
        if not isinstance(bill, dict):
            continue

        bill["sourceFile"] = filename
        clean_bills.append(bill)

    return clean_bills


def is_valid_bill_row(bill: Dict[str, Any]) -> bool:
    """
    Return True only if a parsed bill row has at least one useful value.

    This prevents blank parser rows like:
    PDF Bill | 0 kWh | 0 kVA | ₹0
    from going to the dashboard.
    """
    if not isinstance(bill, dict):
        return False

    return (
        float(bill.get("kWh") or 0) > 0
        or float(bill.get("kVAh") or 0) > 0
        or float(bill.get("totalBillAmount") or 0) > 0
        or float(bill.get("actualDemandKVA") or 0) > 0
        or float(bill.get("billingDemandKVA") or 0) > 0
    )


@app.post("/api/extract-multiple")
async def extract_multiple(files: list[UploadFile] = File(...)) -> Dict[str, Any]:
    parsed_bills: List[Dict[str, Any]] = []
    connected_loads: List[Dict[str, Any]] = []
    file_results: List[Dict[str, Any]] = []

    for uploaded_file in files:
        temp_path = save_upload_to_temp(uploaded_file)

        try:
            filename = uploaded_file.filename or "uploaded_file"
            extension = get_file_extension(filename)
            file_type = classify_file(filename)

            if is_bill_file_extension(extension):
                bills = attach_source_file_to_bills(
                    parse_bill_safely(temp_path),
                    filename,
                )

                # Keep only usable bill rows.
                # This removes zero rows caused by scanned PDFs/OCR failure.
                valid_bills = [bill for bill in bills if is_valid_bill_row(bill)]

                if valid_bills:
                    parsed_bills.extend(valid_bills)

                    file_results.append({
                        "fileName": filename,
                        "fileType": "BILL_FILE",
                        "status": "success",
                        "rowsExtracted": len(valid_bills),
                        "message": f"Extracted {len(valid_bills)} valid bill row(s).",
                    })
                else:
                    file_results.append({
                        "fileName": filename,
                        "fileType": "BILL_FILE",
                        "status": "failed",
                        "rowsExtracted": 0,
                        "message": "No usable bill data could be extracted from this file.",
                    })

            elif extension in [".xlsx", ".xls"]:
                loads = parse_connected_load_safely(temp_path, filename)

                if loads:
                    connected_loads.extend(loads)

                    file_results.append({
                        "fileName": filename,
                        "fileType": "CONNECTED_LOAD_FILE",
                        "status": "success",
                        "rowsExtracted": len(loads),
                        "message": f"Extracted {len(loads)} connected load row(s).",
                    })
                else:
                    bills = attach_source_file_to_bills(
                        parse_excel_safely(temp_path, filename),
                        filename,
                    )

                    # Keep only usable Excel bill rows.
                    valid_bills = [bill for bill in bills if is_valid_bill_row(bill)]

                    if valid_bills:
                        parsed_bills.extend(valid_bills)

                        file_results.append({
                            "fileName": filename,
                            "fileType": "EXCEL_BILL_FILE",
                            "status": "success",
                            "rowsExtracted": len(valid_bills),
                            "message": f"Extracted {len(valid_bills)} valid Excel bill row(s).",
                        })
                    else:
                        file_results.append({
                            "fileName": filename,
                            "fileType": file_type,
                            "status": "failed",
                            "rowsExtracted": 0,
                            "message": "No connected load or usable bill data could be extracted from this Excel file.",
                        })

            elif extension == ".csv":
                bills = attach_source_file_to_bills(
                    parse_csv_safely(temp_path, filename),
                    filename,
                )

                # Keep only usable CSV bill rows.
                valid_bills = [bill for bill in bills if is_valid_bill_row(bill)]

                if valid_bills:
                    parsed_bills.extend(valid_bills)

                    file_results.append({
                        "fileName": filename,
                        "fileType": "CSV_FILE",
                        "status": "success",
                        "rowsExtracted": len(valid_bills),
                        "message": f"Extracted {len(valid_bills)} valid CSV bill row(s).",
                    })
                else:
                    file_results.append({
                        "fileName": filename,
                        "fileType": "CSV_FILE",
                        "status": "failed",
                        "rowsExtracted": 0,
                        "message": "No usable bill data could be extracted from this CSV file.",
                    })

            else:
                file_results.append({
                    "fileName": filename,
                    "fileType": "UNKNOWN_FILE",
                    "status": "failed",
                    "rowsExtracted": 0,
                    "message": "Unsupported file type.",
                })

        except Exception as error:
            file_results.append({
                "fileName": uploaded_file.filename or "uploaded_file",
                "fileType": classify_file(uploaded_file.filename or ""),
                "status": "failed",
                "rowsExtracted": 0,
                "message": str(error),
            })

        finally:
            safe_remove(temp_path)

    successful_files = len(
        [item for item in file_results if item["status"] == "success"]
    )

    failed_files = len(
        [item for item in file_results if item["status"] == "failed"]
    )

    return {
        "status": "completed",
        "bills": parsed_bills,
        "connectedLoads": connected_loads,
        "fileResults": file_results,
        "summary": {
            "totalFiles": len(files),
            "billRows": len(parsed_bills),
            "connectedLoadRows": len(connected_loads),
            "successfulFiles": successful_files,
            "failedFiles": failed_files,
        },
    }